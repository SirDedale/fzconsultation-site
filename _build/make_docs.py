"""Generate the downloadable DOCX templates (FR + EN) into assets/downloads/.
Run by build.py. Requires python-docx."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from tools import RANSOM, RTORPO, AIPOL, AISCORE, TEMPLATE_DISCLAIMER

BRAND = RGBColor(0x00, 0x52, 0xB2)
INK = RGBColor(0x16, 0x22, 0x2E)
MUTED = RGBColor(0x56, 0x65, 0x6F)

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def set_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in t.rows:
        for c, w in zip(row.cells, widths): c.width = Inches(w)
    for col, w in zip(t.columns, widths): col.width = Inches(w)

def base(landscape=False, lang="en"):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin"): setattr(s, side, Inches(0.8 if not landscape else 0.6))
    s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.2
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 12)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = BRAND
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        h.paragraph_format.space_before = Pt(14 if name != "Title" else 0); h.paragraph_format.space_after = Pt(6)
    d.core_properties.language = lang
    d.core_properties.author = "FZ Consultation"
    # footer
    fp = s.footer.paragraphs[0]
    fp.text = ("FZ Consultation · fzconsultation.xyz · " + ("Modèle général, ne constitue pas un avis juridique" if lang == "fr" else "General template, not legal advice"))
    fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = MUTED
    return d

def intro(d, title, lede, lang):
    d.add_paragraph(title, style="Title")
    p = d.add_paragraph(lede); p.runs[0].font.size = Pt(11.5); p.runs[0].font.color.rgb = MUTED
    note = d.add_paragraph(TEMPLATE_DISCLAIMER[0 if lang == "fr" else 1])
    note.runs[0].font.size = Pt(9); note.runs[0].italic = True; note.runs[0].font.color.rgb = MUTED

def checklist(lang):
    i = 0 if lang == "fr" else 1
    d = base(lang=lang)
    intro(d, RANSOM["title"][i], RANSOM["lede"][i], lang)
    d.add_paragraph(RANSOM["how"][i])
    for (title, items) in RANSOM["phases"]:
        d.add_heading(title[i], level=1)
        t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
        for c, txt in zip(t.rows[0].cells, ["", "Action", "Responsable / heure" if lang == "fr" else "Owner / time"]):
            c.text = txt; shade(c, "E6EEF8")
            for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(9)
        for k, it in enumerate(items):
            row = t.add_row().cells
            row[0].text = "☐"; row[1].text = it[i]; row[2].text = ""
            row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_widths(t, [0.4, 4.4, 2.1])
    d.add_heading("Contacts d'urgence" if lang == "fr" else "Emergency contacts", level=1)
    labels = (["Direction / responsable de crise", "Fournisseur d'intervention sur incident", "Assureur cyber (numéro de police)", "Conseiller juridique", "Autorité de cybersécurité / police", "Hébergeur / infonuagique"]
              if lang == "fr" else ["Crisis lead / executive", "Incident response provider", "Cyber insurer (policy number)", "Legal counsel", "Cybersecurity agency / police", "Hosting / cloud provider"])
    t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ["Rôle", "Nom", "Téléphone (hors bande)"] if lang == "fr" else ["Role", "Name", "Phone (out of band)"]):
        c.text = txt; shade(c, "E6EEF8"); c.paragraphs[0].runs[0].bold = True
    for l in labels:
        r = t.add_row().cells; r[0].text = l
    set_widths(t, [2.9, 2.0, 2.0])
    d.add_heading("Sources" if lang == "fr" else "Sources", level=2)
    for name, url in RANSOM["sources"]:
        d.add_paragraph(f"{name} — {url}", style="List Bullet")
    return d

def worksheet(lang, T=RTORPO, widths=None):
    i = 0 if lang == "fr" else 1
    d = base(landscape=True, lang=lang)
    intro(d, T["title"][i], T["lede"][i], lang)
    d.add_paragraph(T["how"][i])
    dt = T.get("defs_title", ("Définitions", "Definitions"))
    d.add_heading(dt[i], level=1)
    for term, desc in T["defs"]:
        p = d.add_paragraph(); r = p.add_run(term[i] + " — "); r.bold = True; p.add_run(desc[i])
    d.add_heading("Tableau" if lang == "fr" else "Worksheet", level=1)
    cols = T["columns"]
    t = d.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"
    for c, col in zip(t.rows[0].cells, cols):
        c.text = col[i]; shade(c, "E6EEF8")
        for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(8.5)
    ex = t.add_row().cells
    for c, v in zip(ex, T["example"]):
        c.text = v[i]
        for r in c.paragraphs[0].runs: r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = MUTED
    for _ in range(12): t.add_row()
    set_widths(t, widths or [1.3, 0.95, 0.8, 0.55, 0.6, 0.6, 0.6, 0.6, 1.8, 2.0])
    cap = d.add_paragraph(("La première ligne est un exemple fictif." if lang == "fr" else "The first row is a fictional example."))
    cap.runs[0].font.size = Pt(8.5); cap.runs[0].italic = True
    d.add_heading("Conseils" if lang == "fr" else "Tips", level=1)
    for h, tx in T["tips"]:
        p = d.add_paragraph(style="List Bullet"); r = p.add_run(h[i] + ". "); r.bold = True; p.add_run(tx[i])
    return d

def policy(lang):
    i = 0 if lang == "fr" else 1
    d = base(lang=lang)
    intro(d, AIPOL["title"][i], AIPOL["lede"][i], lang)
    d.add_paragraph(AIPOL["how"][i])
    meta = [("Organisation", "[Nom de l'organisation]"), ("Version", "[1.0]"), ("Date d'entrée en vigueur", "[date]"), ("Propriétaire", "[fonction]")] if lang == "fr" else [("Organization", "[Organization name]"), ("Version", "[1.0]"), ("Effective date", "[date]"), ("Owner", "[role]")]
    t = d.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in meta:
        r = t.add_row().cells; r[0].text = k; r[1].text = v; shade(r[0], "E6EEF8")
    set_widths(t, [2.0, 4.9])
    for n, (h, tx) in enumerate(AIPOL["sections"], 1):
        d.add_heading(f"{n}. {h[i]}", level=1)
        d.add_paragraph(tx[i])
    return d

def main(out):
    os.makedirs(out, exist_ok=True)
    for lang in ("fr", "en"):
        checklist(lang).save(os.path.join(out, f"{RANSOM['slug']}-{lang}.docx"))
        worksheet(lang).save(os.path.join(out, f"{RTORPO['slug']}-{lang}.docx"))
        policy(lang).save(os.path.join(out, f"{AIPOL['slug']}-{lang}.docx"))
        worksheet(lang, AISCORE, [1.0, 0.65, 1.1, 0.65, 0.7, 0.75, 0.95, 0.95, 0.55, 0.65, 1.1, 0.7]).save(os.path.join(out, f"{AISCORE['slug']}-{lang}.docx"))

if __name__ == "__main__":
    main(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "downloads"))
